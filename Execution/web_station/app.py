"""
DOU Trend Update Web Workstation
Flask web application for appending new Hawi build data to DOU Trend Exec Progress.

Run:
    python Execution/web_station/app.py
Then open: http://localhost:5000
"""
import os
import subprocess
from flask import Flask, render_template, request, jsonify, send_file
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Color
from openpyxl.utils import get_column_letter

app = Flask(__name__)

# ── Config ────────────────────────────────────────────────────────────────────
# app.py lives in Execution/web_station/ → parent is Execution/
EXEC_DIR    = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TGT_FILE    = os.path.join(EXEC_DIR, 'Hawi_GDoUv6_Exec_Progress_Joy.xlsx')
OUTPUT_NAME = 'SM8975_GDoUv6_Power_Breakdown_template_v2 (Output).xlsx'
LOCAL_DIR   = EXEC_DIR
SERVER_BASE = r'\\grilled\ptt_china_la_golden\PT\HAWI'

# ── Shared style helpers ──────────────────────────────────────────────────────
THIN        = Side(border_style='thin')
BORDER_ALL  = Border(left=THIN, right=THIN, top=THIN,  bottom=THIN)
BORDER_HDR  = Border(left=THIN, right=THIN, top=None,  bottom=THIN)
ALIGN_CC    = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_CC_NW = Alignment(horizontal='center', vertical='center', wrap_text=False)

def fill_rgb(hex_rgb):
    return PatternFill(fill_type='solid', fgColor=Color(rgb=hex_rgb))

def fill_theme(theme_idx, tint):
    return PatternFill(fill_type='solid', fgColor=Color(theme=theme_idx, tint=tint))

STYLE_LABEL = {
    'fill':  fill_theme(5, 0.7999816888943144),
    'font':  Font(bold=True, size=20, name='Calibri', color='FF000000'),
    'align': ALIGN_CC,
}
STYLE_SEC_HDR = {
    'fill':   fill_rgb('FF99CC00'),
    'font':   Font(bold=True, size=16, name='Calibri', color='FF000000'),
    'align':  ALIGN_CC,
    'border': BORDER_HDR,
}
STYLE_COL_HDR = {
    'fill':   fill_rgb('FF244062'),
    'font':   Font(bold=True, size=11, name='Calibri', color='FFFFFFFF'),
    'align':  ALIGN_CC,
    'border': BORDER_HDR,
}
STYLE_UC_ID = {
    'font':   Font(bold=True, size=12, name='Calibri'),
    'align':  ALIGN_CC,
    'border': BORDER_ALL,
}
STYLE_DATA = {
    'font':          Font(bold=False, size=12, name='Calibri', color='FF000000'),
    'align':         ALIGN_CC_NW,
    'border':        BORDER_ALL,
    'number_format': '0.0',
}
STYLE_DOU_LABEL = {
    'fill':   fill_theme(4, 0.5999938962981048),
    'font':   Font(bold=True, size=12, name='Calibri'),
    'align':  ALIGN_CC,
    'border': BORDER_ALL,
}
STYLE_DOU_VAL = {
    'fill':          fill_rgb('FFFFFF99'),
    'font':          Font(bold=False, size=12, name='Calibri', color='FF000000'),
    'align':         ALIGN_CC_NW,
    'border':        BORDER_ALL,
    'number_format': '0.00',
}

def apply_style(cell, style):
    if 'fill'          in style: cell.fill          = style['fill']
    if 'font'          in style: cell.font          = style['font']
    if 'align'         in style: cell.alignment     = style['align']
    if 'border'        in style: cell.border        = style['border']
    if 'number_format' in style: cell.number_format = style['number_format']

# ── Routes ────────────────────────────────────────────────────────────────────
@app.route('/')
def index():
    return render_template('index.html')


@app.route('/api/list-builds')
def list_builds():
    """List available build folders on the server."""
    try:
        result = subprocess.run(
            ['cmd', '/c', f'dir "{SERVER_BASE}" /b /ad'],
            capture_output=True, text=True, timeout=15
        )
        builds = [b.strip() for b in result.stdout.strip().split('\n') if b.strip()]
        builds.sort(reverse=True)
        return jsonify({'success': True, 'builds': builds, 'base': SERVER_BASE})
    except subprocess.TimeoutExpired:
        return jsonify({'success': False, 'error': 'Server connection timed out (15 s)'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


@app.route('/api/append-build', methods=['POST'])
def append_build():
    """Copy Output file, read Summary, append to DOU_Trend with formatting."""
    data        = request.json or {}
    build_path  = data.get('build_path', '').strip().strip('"')
    build_label = data.get('build_label', '').strip()

    if not build_path or not build_label:
        return jsonify({'success': False, 'error': 'Missing build_path or build_label'})

    # If the user accidentally included the filename in the path, strip it
    if build_path.lower().endswith(OUTPUT_NAME.lower()):
        build_path = os.path.dirname(build_path)
    # Remove any trailing backslash/slash
    build_path = build_path.rstrip('/\\')

    src_file = os.path.join(LOCAL_DIR, OUTPUT_NAME)
    steps    = []

    # ── Step 1: Copy file ─────────────────────────────────────────────────────
    steps.append({'id': 1, 'name': 'Copy Output file from server', 'status': 'running'})
    try:
        rc = subprocess.run(
            ['robocopy', build_path, LOCAL_DIR, OUTPUT_NAME],
            capture_output=True, text=True, timeout=120
        )
        if not os.path.exists(src_file):
            steps[-1].update(status='error', detail=rc.stderr or rc.stdout or 'File not found after copy')
            return jsonify({'success': False, 'error': 'File copy failed — check the server path.', 'steps': steps})
        steps[-1].update(status='done', detail=f'Saved to {src_file}')
    except subprocess.TimeoutExpired:
        steps[-1].update(status='error', detail='Robocopy timed out after 120 s')
        return jsonify({'success': False, 'error': 'Robocopy timed out', 'steps': steps})
    except Exception as e:
        steps[-1].update(status='error', detail=str(e))
        return jsonify({'success': False, 'error': str(e), 'steps': steps})

    # ── Step 2: Read Summary sheet ────────────────────────────────────────────
    steps.append({'id': 2, 'name': 'Read Summary sheet', 'status': 'running'})
    try:
        src_wb   = load_workbook(src_file, data_only=True)
        src_ws   = src_wb['Summary']
        src_rows = list(src_ws.iter_rows(values_only=True))

        summary = {}
        for row in src_rows[3:]:
            if row[0]:
                summary[row[0]] = row

        measured_count = sum(1 for r in summary.values() if r[3] and r[3] != 0)
        steps[-1].update(status='done',
                         detail=f'{len(summary)} UCs found, {measured_count} with Battery Adjusted data')
    except Exception as e:
        steps[-1].update(status='error', detail=str(e))
        return jsonify({'success': False, 'error': f'Failed to read Summary sheet: {e}', 'steps': steps})

    # ── Step 3: Open target and detect append position ────────────────────────
    steps.append({'id': 3, 'name': 'Detect append position in DOU_Trend', 'status': 'running'})
    try:
        tgt_wb = load_workbook(TGT_FILE)
        tgt_ws = tgt_wb['DOU_Trend']

        last_used_col = 0
        for col in range(1, tgt_ws.max_column + 1):
            if tgt_ws.cell(row=1, column=col).value or tgt_ws.cell(row=2, column=col).value:
                last_used_col = col

        block_num = (last_used_col // 46) + 1
        NEW_START = block_num * 46 + 1
        steps[-1].update(status='done',
                         detail=f'Appending at column {NEW_START} ({get_column_letter(NEW_START)})')
    except Exception as e:
        steps[-1].update(status='error', detail=str(e))
        return jsonify({'success': False, 'error': f'Failed to open target file: {e}', 'steps': steps})

    # ── Step 4: Write data with formatting ────────────────────────────────────
    steps.append({'id': 4, 'name': 'Write data with formatting', 'status': 'running'})
    try:
        # Row 1: build label
        c = tgt_ws.cell(row=1, column=NEW_START + 2, value=f'       {build_label}')
        apply_style(c, STYLE_LABEL)

        # Row 2: column headers (copy from existing block)
        hawi_hdr = [tgt_ws.cell(row=2, column=col).value for col in range(1, 47)]
        for offset, val in enumerate(hawi_hdr):
            c = tgt_ws.cell(row=2, column=NEW_START + offset, value=val)
            apply_style(c, STYLE_SEC_HDR if offset == 0 else STYLE_COL_HDR)

        # Rows 3–31: per-UC data
        for row_idx in range(3, 32):
            uc_id  = tgt_ws.cell(row=row_idx, column=1).value
            weight = tgt_ws.cell(row=row_idx, column=2).value

            c = tgt_ws.cell(row=row_idx, column=NEW_START, value=uc_id)
            apply_style(c, STYLE_UC_ID)

            c = tgt_ws.cell(row=row_idx, column=NEW_START + 1, value=weight)
            apply_style(c, STYLE_UC_ID)
            c.number_format = '0.00%'

            if uc_id in summary:
                src = summary[uc_id]
                c = tgt_ws.cell(row=row_idx, column=NEW_START + 2, value=src[3])
                apply_style(c, STYLE_DATA)
                for k, v in enumerate(src[4:46], start=3):
                    c = tgt_ws.cell(row=row_idx, column=NEW_START + k, value=v)
                    apply_style(c, STYLE_DATA)
            else:
                for k in range(2, 45):
                    c = tgt_ws.cell(row=row_idx, column=NEW_START + k)
                    apply_style(c, STYLE_DATA)

        # Row 32: DoU Current
        measured = []
        for row_idx in range(3, 32):
            uc_id  = tgt_ws.cell(row=row_idx, column=1).value
            weight = tgt_ws.cell(row=row_idx, column=2).value
            batt   = tgt_ws.cell(row=row_idx, column=NEW_START + 2).value
            if uc_id and weight and batt:
                measured.append((row_idx, weight))

        total_weight = sum(w for _, w in measured)

        c = tgt_ws.cell(row=32, column=NEW_START, value='DoU Current')
        apply_style(c, STYLE_DOU_LABEL)

        c = tgt_ws.cell(row=32, column=NEW_START + 1, value=total_weight)
        apply_style(c, STYLE_DOU_LABEL)
        c.number_format = '0.00%'

        for offset in range(2, 45):
            col  = NEW_START + offset
            wsum = sum(weight * (tgt_ws.cell(row=r, column=col).value or 0)
                       for r, weight in measured)
            c = tgt_ws.cell(row=32, column=col, value=wsum)
            apply_style(c, STYLE_DOU_VAL)

        # Column widths
        col_widths = [36.18, 11.82, 10.18] + [13.0] * 42 + [8.0]
        for offset, width in enumerate(col_widths):
            tgt_ws.column_dimensions[get_column_letter(NEW_START + offset)].width = width

        steps[-1].update(status='done',
                         detail=f'{len(measured)} UCs written, total weight: {total_weight:.4f}')
    except Exception as e:
        steps[-1].update(status='error', detail=str(e))
        return jsonify({'success': False, 'error': f'Failed to write data: {e}', 'steps': steps})

    # ── Step 5: Save ──────────────────────────────────────────────────────────
    steps.append({'id': 5, 'name': 'Save Excel file', 'status': 'running'})
    try:
        tgt_wb.save(TGT_FILE)
        steps[-1].update(status='done', detail=TGT_FILE)
    except Exception as e:
        steps[-1].update(status='error', detail=str(e))
        return jsonify({'success': False, 'error': f'Failed to save: {e}', 'steps': steps})

    # ── Step 6: Build comparison + delta tables ───────────────────────────────
    steps.append({'id': 6, 'name': 'Build verification comparison', 'status': 'running'})
    try:
        tgt_wb2       = load_workbook(TGT_FILE, data_only=True)
        ws2           = tgt_wb2['DOU_Trend']
        prev_batt_col = NEW_START - 46 + 2

        # ── Read projection Battery Adjusted (rows 67-93, col 3) ──────────────
        proj_data = {}
        for r in range(67, 94):
            uc_p   = ws2.cell(row=r, column=1).value
            batt_p = ws2.cell(row=r, column=3).value
            if uc_p and uc_p != 'Sum':
                proj_data[uc_p] = batt_p

        # ── Projection DoU = weighted sum over measured UCs ───────────────────
        proj_dou = 0.0
        for r_m, wt_m in measured:
            uc_m     = ws2.cell(row=r_m, column=1).value
            proj_val = proj_data.get(uc_m) or 0
            proj_dou += wt_m * proj_val

        # ── Per-UC comparison rows ────────────────────────────────────────────
        comparison = []
        for r in range(3, 33):
            uc  = ws2.cell(row=r, column=1).value
            wt  = ws2.cell(row=r, column=2).value
            prv = ws2.cell(row=r, column=prev_batt_col).value if prev_batt_col >= 3 else None
            new = ws2.cell(row=r, column=NEW_START + 2).value

            if prv is not None or new is not None:
                delta      = (new - prv) if isinstance(prv, (int, float)) and isinstance(new, (int, float)) else None
                # Projection value: per-UC lookup, or weighted DoU for row 32
                proj_batt  = proj_dou if r == 32 else proj_data.get(str(uc) if uc else '')
                delta_proj = (new - proj_batt) if isinstance(new, (int, float)) and isinstance(proj_batt, (int, float)) else None
                comparison.append({
                    'uc_id':      str(uc or ''),
                    'weight':     wt,
                    'previous':   prv,
                    'new':        new,
                    'delta':      delta,
                    'proj_batt':  proj_batt,
                    'delta_proj': delta_proj,
                    'is_dou':     r == 32,
                })

        batt_new       = ws2.cell(row=32, column=NEW_START + 2).value
        batt_prv       = ws2.cell(row=32, column=prev_batt_col).value if prev_batt_col >= 3 else None
        dou_delta      = (batt_new - batt_prv)  if isinstance(batt_new, (int, float)) and isinstance(batt_prv, (int, float))  else None
        dou_delta_proj = (batt_new - proj_dou)  if isinstance(batt_new, (int, float)) and proj_dou is not None else None

        steps[-1].update(status='done', detail=f'DoU Current: {batt_new:.2f} mA' if isinstance(batt_new, float) else '')

        return jsonify({
            'success': True,
            'steps':   steps,
            'result': {
                'build_label':    build_label,
                'column':         NEW_START,
                'column_letter':  get_column_letter(NEW_START),
                'measured_ucs':   len(measured),
                'total_weight':   total_weight,
                'dou_current':    batt_new,
                'dou_previous':   batt_prv,
                'dou_delta':      dou_delta,
                'proj_dou':       proj_dou,
                'dou_delta_proj': dou_delta_proj,
            },
            'comparison': comparison,
        })
    except Exception as e:
        steps[-1].update(status='error', detail=str(e))
        return jsonify({'success': False, 'error': f'Failed to build comparison: {e}', 'steps': steps})


@app.route('/api/download')
def download():
    """Download the updated Exec Progress Excel file."""
    return send_file(TGT_FILE, as_attachment=True,
                     download_name='Hawi_GDoUv6_Exec_Progress_Joy.xlsx')


if __name__ == '__main__':
    print(f"\n{'='*60}")
    print("  DOU Trend Update Workstation")
    print(f"  Target file : {TGT_FILE}")
    print(f"  Local dir   : {LOCAL_DIR}")
    print(f"{'='*60}")
    print("  Open http://localhost:5000 in your browser\n")
    app.run(debug=False, port=5000, host='0.0.0.0')