"""
Append a new build's data (from Summary sheet of Power Breakdown template)
to the right of the existing data in DOU_Trend sheet of the Exec Progress file,
with full cell formatting matching the existing table style.

Usage:
    python Execution/append_build.py <build_folder_path> <build_label>

Example:
    python Execution/append_build.py "\\grilled\ptt_china_la_golden\PT\HAWI\Hawi.LA.1.0-00510-PERF.AL-3\PowerFtrace" M510_0317
"""
import sys
import os
import subprocess
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Color
from openpyxl.utils import get_column_letter

# ── Config ───────────────────────────────────────────────────────────────────
TGT_FILE    = 'Execution/Hawi_GDoUv6_Exec_Progress_Joy.xlsx'
OUTPUT_NAME = 'SM8975_GDoUv6_Power_Breakdown_template_v2 (Output).xlsx'
LOCAL_DIR   = 'Execution'

# ── Parse arguments or prompt interactively ───────────────────────────────────
if len(sys.argv) == 3:
    BUILD_PATH  = sys.argv[1]
    BUILD_LABEL = sys.argv[2]
elif len(sys.argv) == 1:
    print("=== DOU Trend Update ===\n")
    BUILD_PATH  = input("Enter build PowerFtrace folder path\n"
                        "  (e.g. \\\\grilled\\ptt_china_la_golden\\PT\\HAWI\\"
                        "Hawi.LA.1.0-00510-PERF.AL-3\\PowerFtrace): ").strip().strip('"')
    BUILD_LABEL = input("Enter build label (e.g. M510_0317): ").strip()
else:
    print(__doc__)
    sys.exit(1)

SRC_FILE = os.path.join(LOCAL_DIR, OUTPUT_NAME)

# ── Shared style helpers ──────────────────────────────────────────────────────
THIN = Side(border_style='thin')
BORDER_ALL  = Border(left=THIN, right=THIN, top=THIN,  bottom=THIN)
BORDER_HDR  = Border(left=THIN, right=THIN, top=None,  bottom=THIN)
ALIGN_CC    = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_CC_NW = Alignment(horizontal='center', vertical='center', wrap_text=False)

def fill_rgb(hex_rgb):
    return PatternFill(fill_type='solid', fgColor=Color(rgb=hex_rgb))

def fill_theme(theme_idx, tint):
    return PatternFill(fill_type='solid', fgColor=Color(theme=theme_idx, tint=tint))

# Style definitions (matched from existing table)
STYLE_LABEL = {   # Row 1: build label cell
    'fill':   fill_theme(5, 0.7999816888943144),
    'font':   Font(bold=True, size=20, name='Calibri', color='FF000000'),
    'align':  ALIGN_CC,
}
STYLE_SEC_HDR = {  # Row 2, col +0: section name ("Hawi")
    'fill':   fill_rgb('FF99CC00'),
    'font':   Font(bold=True, size=16, name='Calibri', color='FF000000'),
    'align':  ALIGN_CC,
    'border': BORDER_HDR,
}
STYLE_COL_HDR = {  # Row 2, cols +1 to +44: column headers
    'fill':   fill_rgb('FF244062'),
    'font':   Font(bold=True, size=11, name='Calibri', color='FFFFFFFF'),
    'align':  ALIGN_CC,
    'border': BORDER_HDR,
}
STYLE_UC_ID = {    # Rows 3-31, cols +0 and +1: UC ID and weight
    'font':   Font(bold=True, size=12, name='Calibri'),
    'align':  ALIGN_CC,
    'border': BORDER_ALL,
}
STYLE_DATA = {     # Rows 3-31, cols +2 to +44: numeric data
    'font':   Font(bold=False, size=12, name='Calibri', color='FF000000'),
    'align':  ALIGN_CC_NW,
    'border': BORDER_ALL,
    'number_format': '0.0',
}
STYLE_DOU_LABEL = {  # Row 32, cols +0 and +1: DoU Current label/weight
    'fill':   fill_theme(4, 0.5999938962981048),
    'font':   Font(bold=True, size=12, name='Calibri'),
    'align':  ALIGN_CC,
    'border': BORDER_ALL,
}
STYLE_DOU_VAL = {    # Row 32, cols +2 to +44: DoU Current values
    'fill':   fill_rgb('FFFFFF99'),
    'font':   Font(bold=False, size=12, name='Calibri', color='FF000000'),
    'align':  ALIGN_CC_NW,
    'border': BORDER_ALL,
    'number_format': '0.00',
}

def apply_style(cell, style):
    if 'fill'          in style: cell.fill          = style['fill']
    if 'font'          in style: cell.font          = style['font']
    if 'align'         in style: cell.alignment     = style['align']
    if 'border'        in style: cell.border        = style['border']
    if 'number_format' in style: cell.number_format = style['number_format']

# ── Step 1: Copy file from server ─────────────────────────────────────────────
print(f"\n[Step 1] Copying '{OUTPUT_NAME}' from:\n  {BUILD_PATH}")
subprocess.run(['robocopy', BUILD_PATH, LOCAL_DIR, OUTPUT_NAME],
               capture_output=True, text=True)
if not os.path.exists(SRC_FILE):
    print("ERROR: File copy failed. Check the path and try again.")
    sys.exit(1)
print(f"  -> Copied to {SRC_FILE}")

# ── Step 2: Read Summary sheet ────────────────────────────────────────────────
print(f"\n[Step 2] Reading Summary sheet...")
src_wb   = load_workbook(SRC_FILE, data_only=True)
src_ws   = src_wb['Summary']
src_rows = list(src_ws.iter_rows(values_only=True))

summary = {}
for row in src_rows[3:]:
    if row[0]:
        summary[row[0]] = row

measured_count = sum(1 for r in summary.values() if r[3] and r[3] != 0)
print(f"  -> {len(summary)} UCs found, {measured_count} with Battery Adjusted data")

# ── Step 3: Open target and detect append position ────────────────────────────
print(f"\n[Step 3] Opening {TGT_FILE}...")
tgt_wb = load_workbook(TGT_FILE)
tgt_ws = tgt_wb['DOU_Trend']

last_used_col = 0
for col in range(1, tgt_ws.max_column + 1):
    if tgt_ws.cell(row=1, column=col).value or tgt_ws.cell(row=2, column=col).value:
        last_used_col = col

block_num = (last_used_col // 46) + 1
NEW_START = block_num * 46 + 1
print(f"  -> Appending at column {NEW_START} ({get_column_letter(NEW_START)})")

# ── Step 4: Write data + formatting ──────────────────────────────────────────
print(f"\n[Step 4] Writing build '{BUILD_LABEL}' data with formatting...")

# --- Row 1: build label ---
c = tgt_ws.cell(row=1, column=NEW_START + 2, value=f'       {BUILD_LABEL}')
apply_style(c, STYLE_LABEL)

# --- Row 2: headers ---
hawi_hdr = [tgt_ws.cell(row=2, column=col).value for col in range(1, 47)]
for offset, val in enumerate(hawi_hdr):
    c = tgt_ws.cell(row=2, column=NEW_START + offset, value=val)
    apply_style(c, STYLE_SEC_HDR if offset == 0 else STYLE_COL_HDR)

# --- Rows 3-31: per-UC data ---
for row_idx in range(3, 32):
    uc_id  = tgt_ws.cell(row=row_idx, column=1).value
    weight = tgt_ws.cell(row=row_idx, column=2).value

    c = tgt_ws.cell(row=row_idx, column=NEW_START, value=uc_id)
    apply_style(c, STYLE_UC_ID)

    c = tgt_ws.cell(row=row_idx, column=NEW_START + 1, value=weight)
    apply_style(c, STYLE_UC_ID)
    c.number_format = '0.00%'

    if uc_id in summary:
        src = summary[uc_id]
        batt_adj = src[3]
        c = tgt_ws.cell(row=row_idx, column=NEW_START + 2, value=batt_adj)
        apply_style(c, STYLE_DATA)
        for k, v in enumerate(src[4:46], start=3):
            c = tgt_ws.cell(row=row_idx, column=NEW_START + k, value=v)
            apply_style(c, STYLE_DATA)
    else:
        for k in range(2, 45):
            c = tgt_ws.cell(row=row_idx, column=NEW_START + k)
            apply_style(c, STYLE_DATA)

# --- Row 32: DoU Current ---
measured = []
for row_idx in range(3, 32):
    uc_id  = tgt_ws.cell(row=row_idx, column=1).value
    weight = tgt_ws.cell(row=row_idx, column=2).value
    batt   = tgt_ws.cell(row=row_idx, column=NEW_START + 2).value
    if uc_id and weight and batt:
        measured.append((row_idx, weight))

total_weight = sum(w for _, w in measured)

c = tgt_ws.cell(row=32, column=NEW_START, value='DoU Current')
apply_style(c, STYLE_DOU_LABEL)

c = tgt_ws.cell(row=32, column=NEW_START + 1, value=total_weight)
apply_style(c, STYLE_DOU_LABEL)
c.number_format = '0.00%'

for offset in range(2, 45):
    col  = NEW_START + offset
    wsum = sum(weight * (tgt_ws.cell(row=r, column=col).value or 0)
               for r, weight in measured)
    c = tgt_ws.cell(row=32, column=col, value=wsum)
    apply_style(c, STYLE_DOU_VAL)

# --- Column widths for new block ---
col_widths = [36.18, 11.82, 10.18] + [13.0] * 42 + [8.0]
for offset, width in enumerate(col_widths):
    tgt_ws.column_dimensions[get_column_letter(NEW_START + offset)].width = width

# ── Step 5: Save ──────────────────────────────────────────────────────────────
tgt_wb.save(TGT_FILE)
print(f"  -> Saved: {TGT_FILE}")

# ── Step 6: Print comparison summary ─────────────────────────────────────────
print(f"\n=== Done ===")
print(f"Build '{BUILD_LABEL}' appended at column {get_column_letter(NEW_START)} (col {NEW_START}).")
print(f"Measured UCs: {len(measured)}, total weight: {total_weight:.3f}")

tgt_wb2 = load_workbook(TGT_FILE, data_only=True)
ws2 = tgt_wb2['DOU_Trend']
batt_new = ws2.cell(row=32, column=NEW_START + 2).value
print(f"Weighted Battery Adjusted ({BUILD_LABEL}): {batt_new:.2f} mA\n")

prev_batt_col = NEW_START - 46 + 2
print("Per-UC Battery Adjusted comparison:")
print("  {:20s} | {:5s} | {:>12} | {:>12} | {:>8}".format(
    "UC ID", "Wt", "Previous", BUILD_LABEL, "Delta"))
print("  " + "-" * 68)
for r in range(3, 33):
    uc  = ws2.cell(row=r, column=1).value
    wt  = ws2.cell(row=r, column=2).value
    prv = ws2.cell(row=r, column=prev_batt_col).value if prev_batt_col >= 3 else None
    new = ws2.cell(row=r, column=NEW_START + 2).value
    if prv or new:
        prv_s = "{:.2f}".format(prv) if isinstance(prv, (int, float)) else str(prv or '-')
        new_s = "{:.2f}".format(new) if isinstance(new, (int, float)) else str(new or '-')
        delta = "{:+.2f}".format(new - prv) if isinstance(prv, (int, float)) and isinstance(new, (int, float)) else '-'
        print("  {:20s} | {:5s} | {:>12} | {:>12} | {:>8}".format(
            str(uc), str(wt), prv_s, new_s, delta))