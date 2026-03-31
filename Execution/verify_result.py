import openpyxl
wb = openpyxl.load_workbook('Execution/Hawi_GDoUv6_Exec_Progress_Joy.xlsx', data_only=True)
ws = wb['DOU_Trend']
print('DOU_Trend sheet now has', ws.max_column, 'columns (was 46)')
print()
print('Build labels in Row 1:')
print('  Col C  (M446_AU205 section):', ws.cell(row=1, column=3).value)
print('  Col AW (M446_0226 section): ', ws.cell(row=1, column=49).value)
print()
print('DoU Current (Row 32) - Battery Adjusted:')
print('  M446_AU205 (col C):  110.39 mA  [formula-based, reads as None via openpyxl]')
new_dou = ws.cell(row=32, column=49).value
print('  M446_0226  (col AW):', round(new_dou, 2), 'mA')
print()
print('Per-UC Battery Adjusted: M446_AU205 vs M446_0226')
print('  {:20s} | {:5s} | {:>12} | {:>12} | {:>10}'.format('UC ID','Wt','M446_AU205','M446_0226','Delta'))
print('  ' + '-'*70)
for r in range(3, 32):
    uc  = ws.cell(row=r, column=1).value
    wt  = ws.cell(row=r, column=2).value
    old = ws.cell(row=r, column=3).value
    new = ws.cell(row=r, column=49).value
    if old or new:
        old_s = '{:.2f}'.format(old) if old else '-'
        new_s = '{:.2f}'.format(new) if new else '-'
        delta = '{:+.2f}'.format(new - old) if (old and new) else '-'
        print('  {:20s} | {:5s} | {:>12} | {:>12} | {:>10}'.format(str(uc), str(wt), old_s, new_s, delta))